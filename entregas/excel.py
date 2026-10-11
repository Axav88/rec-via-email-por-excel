"""Lectura tolerante del Excel de pedidos."""
from __future__ import annotations

import shutil
import tempfile
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .core import SIN_OC, Fila

FORMATOS_FECHA = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y")


def _norm(texto) -> str:
    """'  Descripción ' -> 'descripcion' (sin acentos, minúsculas, espacios simples)."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(t.lower().split())


def _fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            return from_excel(valor).date()
        except Exception:
            return None
    if isinstance(valor, str):
        for fmt in FORMATOS_FECHA:
            try:
                return datetime.strptime(valor.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _oc(valor) -> str:
    if valor is None or str(valor).strip() == "":
        return SIN_OC
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _piezas(valor) -> int:
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return 0


def leer_filas(ruta: Path) -> tuple[list[Fila], list[str]]:
    """Devuelve (filas, advertencias). Busca el encabezado solo, así que da igual si
    cambian de lugar las columnas; lo único que importa es que se llamen ORDEN DE COMPRA,
    CLIENTE y ENTREGA.
    """
    ruta = Path(ruta)
    if not ruta.exists():
        raise FileNotFoundError(f"No encuentro el Excel: {ruta}")

    filas: list[Fila] = []
    avisos: list[str] = []

    # Se trabaja sobre una copia: así funciona aunque el Excel esté abierto.
    with tempfile.TemporaryDirectory() as tmp:
        copia = Path(tmp) / "copia.xlsx"
        shutil.copy2(ruta, copia)
        wb = load_workbook(copia, read_only=True, data_only=True)
        try:
            for hoja in wb.worksheets:
                filas_hoja = list(hoja.iter_rows(values_only=True))
                cols, inicio = None, 0
                for i, fila in enumerate(filas_hoja[:15]):
                    nombres = {_norm(c): j for j, c in reversed(list(enumerate(fila))) if c is not None}
                    if "entrega" in nombres and "cliente" in nombres:
                        cols, inicio = nombres, i + 1
                        break
                if cols is None:
                    continue

                def celda(fila, nombre):
                    j = cols.get(nombre)
                    return fila[j] if j is not None and j < len(fila) else None

                for n, fila in enumerate(filas_hoja[inicio:], start=inicio + 1):
                    bruto = celda(fila, "entrega")
                    if bruto is None or str(bruto).strip() == "":
                        continue
                    fecha = _fecha(bruto)
                    if fecha is None:
                        avisos.append(f"{hoja.title}, fila {n}: fecha de ENTREGA ilegible ({bruto!r})")
                        continue
                    filas.append(
                        Fila(
                            oc=_oc(celda(fila, "orden de compra")),
                            cliente=str(celda(fila, "cliente") or "(sin cliente)").strip(),
                            fecha=fecha,
                            piezas=_piezas(celda(fila, "cant/pzs")),
                            descripcion=str(celda(fila, "descripcion") or "").strip(),
                        )
                    )
        finally:
            wb.close()

    if not filas and not avisos:
        avisos.append("No encontré filas con ENTREGA. ¿Cambió el nombre de las columnas?")
    return filas, avisos
